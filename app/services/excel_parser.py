from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.services.spreadsheet_metrics import calculate_signal_columns

HYPERLINK_RE = re.compile(
    r'=HYPERLINK\("(?P<url>[^"]+)"\s*,\s*"(?P<label>[^"]+)"\)',
    re.IGNORECASE,
)
URL_IDS_RE = re.compile(r"/tournaments/(?P<tournament_id>\d+)/(?P<match_id>\d+)")
PLAYER_RE = re.compile(
    r"^\s*(?:\((?P<rating1>\d+)\)\s*)?(?P<player1>.+?)\s+vs\s+"
    r"(?:\((?P<rating2>\d+)\)\s*)?(?P<player2>.+?)\s*$",
    re.IGNORECASE,
)
DATE_RE = re.compile(r"(?P<day>\d{2})\.(?P<month>\d{2})\.(?P<year>\d{4})")


@dataclass(slots=True)
class ParsedMatch:
    external_match_id: int
    external_tournament_id: int
    source_url: str
    tournament_date: str
    tournament_name: str
    match_time: str
    match_start_at: datetime
    player_1: str
    player_2: str
    player_1_rating: int | None
    player_2_rating: int | None
    score: str | None
    raw_data: dict[str, Any]


@dataclass(slots=True)
class ParseResult:
    sheet_name: str
    total_rows: int
    matches: list[ParsedMatch]
    warnings: list[str]


def _json_safe(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


def _parse_hyperlink_formula(formula: str) -> tuple[str, str] | None:
    match = HYPERLINK_RE.match(formula.strip())
    if not match:
        return None
    return match.group("url"), match.group("label")


def parse_tournaments_file(path: Path, timezone: str = "Europe/Moscow") -> ParseResult:
    # Первая книга нужна для HYPERLINK и исходных формул, вторая — для
    # сохранённых Excel значений вычисляемых столбцов (EG, EH, CS, CT и т. д.).
    formula_book = load_workbook(path, data_only=False, read_only=True)
    values_book = load_workbook(path, data_only=True, read_only=True)
    if "Турниры" not in formula_book.sheetnames:
        formula_book.close()
        values_book.close()
        raise ValueError('В файле отсутствует обязательный лист "Турниры".')

    formula_sheet = formula_book["Турниры"]
    values_sheet = values_book["Турниры"]
    current_date: str | None = None
    current_headers: list[str] = []
    current_tournament_name = "Турнир"
    parsed: list[ParsedMatch] = []
    warnings: list[str] = []
    tz = ZoneInfo(timezone)

    formula_rows = formula_sheet.iter_rows()
    values_rows = values_sheet.iter_rows()
    for row_number, (formula_row, values_row) in enumerate(zip(formula_rows, values_rows), start=1):
        first_formula = formula_row[0].value if formula_row else None
        first_value = values_row[0].value if values_row else None
        date_match = DATE_RE.search(str(first_value)) if first_value else None
        if date_match:
            current_date = date_match.group(0)
            current_headers = [
                str(cell.value).strip() if cell.value is not None else f"COL_{index + 1}"
                for index, cell in enumerate(values_row)
            ]
            continue

        # Первая строка турнира содержит HYPERLINK в A и название лиги.
        if isinstance(first_formula, str) and first_formula.upper().startswith("=HYPERLINK"):
            tournament_link = _parse_hyperlink_formula(first_formula)
            if tournament_link is not None:
                current_tournament_name = tournament_link[1].strip()

        if not current_date or len(formula_row) < 3:
            continue

        time_value = values_row[1].value
        formula = formula_row[2].value
        if not time_value or not isinstance(formula, str) or not formula.upper().startswith("=HYPERLINK"):
            continue

        hyperlink = _parse_hyperlink_formula(formula)
        if hyperlink is None:
            warnings.append(f"Строка {row_number}: не удалось разобрать HYPERLINK.")
            continue
        source_url, label = hyperlink

        ids_match = URL_IDS_RE.search(source_url)
        players_match = PLAYER_RE.match(label)
        if ids_match is None or players_match is None:
            warnings.append(f"Строка {row_number}: не удалось определить ID или игроков.")
            continue

        match_time = str(time_value).strip()[:5]
        try:
            start_at = datetime.strptime(f"{current_date} {match_time}", "%d.%m.%Y %H:%M").replace(tzinfo=tz)
        except ValueError:
            warnings.append(f"Строка {row_number}: неверные дата/время {current_date} {match_time}.")
            continue

        raw_data: dict[str, Any] = {
            "_row_number": row_number,
            "_tournament_name": current_tournament_name,
        }
        for index, (formula_cell, value_cell) in enumerate(zip(formula_row, values_row)):
            column = get_column_letter(index + 1)
            header = current_headers[index] if index < len(current_headers) else f"COL_{index + 1}"
            value = value_cell.value
            formula_value = formula_cell.value
            # Буква столбца — основной стабильный ключ для правил.
            raw_data[column] = _json_safe(value)
            raw_data[f"{header}__{index + 1}"] = _json_safe(value)
            if isinstance(formula_value, str) and formula_value.startswith("="):
                raw_data[f"_{column}_formula"] = formula_value

        calculate_signal_columns(raw_data)

        parsed.append(
            ParsedMatch(
                external_match_id=int(ids_match.group("match_id")),
                external_tournament_id=int(ids_match.group("tournament_id")),
                source_url=source_url,
                tournament_date=current_date,
                tournament_name=current_tournament_name,
                match_time=match_time,
                match_start_at=start_at,
                player_1=players_match.group("player1").strip(),
                player_2=players_match.group("player2").strip(),
                player_1_rating=int(players_match.group("rating1")) if players_match.group("rating1") else None,
                player_2_rating=int(players_match.group("rating2")) if players_match.group("rating2") else None,
                score=str(values_row[7].value).strip() if len(values_row) > 7 and values_row[7].value is not None else None,
                raw_data=raw_data,
            )
        )

    formula_book.close()
    values_book.close()
    return ParseResult(
        sheet_name=formula_sheet.title,
        total_rows=formula_sheet.max_row,
        matches=parsed,
        warnings=warnings,
    )
